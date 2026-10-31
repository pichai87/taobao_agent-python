"""A separate, enriched UCI import for approved range and dimension analysis."""

import argparse
import json
import sqlite3
from contextlib import closing
from dataclasses import asdict, dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from taobao_agent.metric_catalog import load_metric_catalog
from taobao_agent.metric_search import MetricDefinition
from taobao_agent.metric_sql import AggregateQuery, build_aggregate_query, validate_aggregate_query
from taobao_agent.uci_online_retail import (
    SOURCE_ID, UCIDateOutOfRange, UCIDataProfile, UCIQueryError, _SOURCE_COLUMNS,
    _line_record, _source_hash,
)


ENRICHED_SOURCE_ID = SOURCE_ID + "-enriched-v1"
ENRICHED_CATALOG_PATH = (
    Path(__file__).resolve().parents[2] / "examples" / "uci_online_retail_enriched_metrics.json"
)
_APPROVED_COLUMNS = frozenset({
    "invoice_date", "record_type", "line_amount_pence", "quantity", "line_count",
    "country", "stock_code",
})


@dataclass(frozen=True)
class EnrichedImportSummary:
    source_sha256: str
    positive_rows: int
    cancelled_rows: int
    other_rows: int
    first_date: str
    last_date: str


class GroupValue(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    group: str
    value: int = Field(ge=0)


class AggregateResult(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    metric_name: str
    start: date
    end: date
    group_by: str | None
    total: int = Field(ge=0)
    groups: tuple[GroupValue, ...] = ()
    unit: str
    source: str = "uci"

    @model_validator(mode="after")
    def check_groups(self) -> "AggregateResult":
        if self.start > self.end or (self.group_by is None and self.groups):
            raise ValueError("聚合结果日期或分组无效")
        if self.group_by is not None and sum(item.value for item in self.groups) != self.total:
            raise ValueError("分组合计与总额不一致")
        if len({item.group for item in self.groups}) != len(self.groups):
            raise ValueError("分组名称重复")
        return self


def import_uci_enriched(source_path: Path, database_path: Path) -> EnrichedImportSummary:
    """Create a new SQLite file; retain country and product code, never customer ID."""

    if database_path.exists():
        raise FileExistsError(f"目标数据库已存在，不会覆盖：{database_path}")
    source_hash = _source_hash(source_path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        if tuple(next(rows, ())) != _SOURCE_COLUMNS:
            raise ValueError("UCI 文件表头与 Online Retail 官方字段不一致")
        database_path.parent.mkdir(parents=True, exist_ok=True)
        with closing(sqlite3.connect(database_path)) as connection:
            try:
                connection.executescript(
                    "CREATE TABLE retail_lines (invoice_date TEXT NOT NULL, "
                    "record_type TEXT NOT NULL CHECK (record_type IN ('positive', 'cancel', 'other')), "
                    "line_amount_pence INTEGER NOT NULL, quantity INTEGER NOT NULL, "
                    "line_count INTEGER NOT NULL, country TEXT NOT NULL, stock_code TEXT NOT NULL);"
                    "CREATE TABLE import_metadata (source_id TEXT NOT NULL, source_sha256 TEXT NOT NULL, "
                    "first_date TEXT NOT NULL, last_date TEXT NOT NULL);"
                )
                counts = {"positive": 0, "cancel": 0, "other": 0}
                first: str | None = None
                last: str | None = None
                batch: list[tuple[str, str, int, int, int, str, str]] = []
                for row_number, row in enumerate(rows, start=2):
                    day, kind, pence = _line_record(row, row_number)
                    quantity, stock_code, country = row[3], row[1], row[7]
                    if not isinstance(quantity, int) or isinstance(quantity, bool):
                        raise ValueError(f"UCI 第 {row_number} 行的件数不是整数")
                    if not isinstance(country, str) or not country.strip():
                        raise ValueError(f"UCI 第 {row_number} 行的国家为空")
                    if not isinstance(stock_code, (str, int)) or not str(stock_code).strip():
                        raise ValueError(f"UCI 第 {row_number} 行的商品编码为空")
                    batch.append((day, kind, pence, quantity, 1, country.strip(), str(stock_code).strip()))
                    counts[kind] += 1
                    first = min(first, day) if first else day
                    last = max(last, day) if last else day
                    if len(batch) == 5000:
                        connection.executemany("INSERT INTO retail_lines VALUES (?, ?, ?, ?, ?, ?, ?)", batch)
                        batch.clear()
                if batch:
                    connection.executemany("INSERT INTO retail_lines VALUES (?, ?, ?, ?, ?, ?, ?)", batch)
                if first is None or last is None:
                    raise ValueError("UCI 文件没有可导入的交易明细")
                connection.execute("CREATE INDEX retail_lines_date_type ON retail_lines(invoice_date, record_type)")
                connection.execute("CREATE INDEX retail_lines_country ON retail_lines(country)")
                connection.execute("CREATE INDEX retail_lines_stock_code ON retail_lines(stock_code)")
                connection.execute(
                    "INSERT INTO import_metadata VALUES (?, ?, ?, ?)",
                    (ENRICHED_SOURCE_ID, source_hash, first, last),
                )
                connection.commit()
            except BaseException:
                connection.rollback()
                raise
    except BaseException:
        database_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return EnrichedImportSummary(source_hash, counts["positive"], counts["cancel"],
                                 counts["other"], first, last)


def source_catalog(database_path: Path) -> Sequence[MetricDefinition]:
    """Select the compatible built-in catalog from the imported source marker."""

    from taobao_agent.uci_online_retail import UCI_CATALOG_PATH

    try:
        uri = f"{database_path.resolve().as_uri()}?mode=ro"
        with closing(sqlite3.connect(uri, uri=True)) as connection:
            row = connection.execute("SELECT source_id FROM import_metadata").fetchone()
    except (OSError, sqlite3.Error, ValueError) as exc:
        raise UCIQueryError("无法识别本地 UCI 数据源") from exc
    if row == (ENRICHED_SOURCE_ID,):
        return load_metric_catalog(ENRICHED_CATALOG_PATH)
    if row == (SOURCE_ID,):
        return load_metric_catalog(UCI_CATALOG_PATH)
    raise UCIQueryError("数据库不是本项目的 UCI 导入结果")


def is_enriched_catalog(catalog: Sequence[MetricDefinition]) -> bool:
    return tuple(catalog) == tuple(load_metric_catalog(ENRICHED_CATALOG_PATH))


def inspect_enriched_metric(metric: MetricDefinition, database_path: Path) -> UCIDataProfile:
    """Verify a catalog metric against the imported enhanced table before describing it."""

    if metric not in load_metric_catalog(ENRICHED_CATALOG_PATH):
        raise UCIQueryError("指标不在增强版 UCI 数据批准范围内")
    try:
        uri = f"{database_path.resolve().as_uri()}?mode=ro"
        with closing(sqlite3.connect(uri, uri=True)) as connection:
            connection.execute("PRAGMA query_only = ON")
            row = connection.execute(
                "SELECT source_id, first_date, last_date FROM import_metadata"
            ).fetchone()
            if row is None or row[0] != ENRICHED_SOURCE_ID:
                raise UCIQueryError("数据库不是增强版 UCI 导入结果")
            actual = frozenset(item[1] for item in connection.execute('PRAGMA table_info("retail_lines")'))
            if not _APPROVED_COLUMNS <= actual:
                raise UCIQueryError("增强版数据表缺少批准的字段")
            return UCIDataProfile(row[1], row[2], tuple(sorted(_APPROVED_COLUMNS)))
    except (OSError, sqlite3.Error, ValueError) as exc:
        raise UCIQueryError("增强版 UCI 元信息检查失败") from exc


def query_uci_aggregate(
    metric: MetricDefinition, start: date, end: date, database_path: Path,
    group_by: str | None = None, *, candidate: AggregateQuery | None = None,
) -> AggregateResult:
    """Validate the complete query against source, catalog and actual columns."""

    if metric not in load_metric_catalog(ENRICHED_CATALOG_PATH):
        raise UCIQueryError("指标不在增强版 UCI 数据批准范围内")
    if start > end or end - start > timedelta(days=30):
        raise UCIQueryError("日期范围最多 31 天")
    if group_by not in (None, "国家", "商品编码"):
        raise UCIQueryError("分组维度未获批准")
    try:
        uri = f"{database_path.resolve().as_uri()}?mode=ro"
        with closing(sqlite3.connect(uri, uri=True)) as connection:
            connection.execute("PRAGMA query_only = ON")
            metadata = connection.execute(
                "SELECT source_id, first_date, last_date FROM import_metadata"
            ).fetchone()
            if metadata is None or metadata[0] != ENRICHED_SOURCE_ID:
                raise UCIQueryError("数据库不是增强版 UCI 导入结果")
            if start.isoformat() < metadata[1] or end.isoformat() > metadata[2]:
                raise UCIDateOutOfRange(metadata[1], metadata[2])
            actual = frozenset(row[1] for row in connection.execute('PRAGMA table_info("retail_lines")'))
            approved = {"retail_lines": actual & _APPROVED_COLUMNS}
            query = candidate or build_aggregate_query(metric, start, end, group_by)
            validated = validate_aggregate_query(query, metric, start, end, group_by, approved)
            rows = connection.execute(validated.statement, validated.parameters).fetchall()
            if group_by is None:
                if len(rows) != 1 or not isinstance(rows[0][0], int):
                    raise UCIQueryError("聚合结果格式无效")
                return AggregateResult(metric_name=metric.name, start=start, end=end,
                                       group_by=None, total=rows[0][0], unit=metric.calculation.unit)
            if len(rows) > 5000 or any(
                not isinstance(group, str) or not isinstance(value, int) for group, value in rows
            ):
                raise UCIQueryError("分组结果超过上限或格式无效")
            groups = tuple(GroupValue(group=group, value=value) for group, value in rows)
            return AggregateResult(metric_name=metric.name, start=start, end=end,
                                   group_by=group_by, total=sum(item.value for item in groups),
                                   groups=groups, unit=metric.calculation.unit)
    except (OSError, sqlite3.Error, ValueError) as exc:
        raise UCIQueryError("增强版 UCI 查询未通过校验或执行失败") from exc


def main() -> None:
    parser = argparse.ArgumentParser(description="从 UCI Excel 新建含国家和商品编码的 SQLite")
    parser.add_argument("source", type=Path)
    parser.add_argument("database", type=Path, help="新文件路径；已有文件不会覆盖")
    args = parser.parse_args()
    try:
        summary = import_uci_enriched(args.source, args.database)
    except (OSError, ValueError, sqlite3.Error) as exc:
        parser.exit(2, f"导入失败：{exc}\n")
    print(json.dumps(asdict(summary), ensure_ascii=False))
