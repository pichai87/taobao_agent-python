"""Import UCI Online Retail lines and query the approved local SQLite copy."""

import argparse
import hashlib
import json
import sqlite3
from contextlib import closing
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

from taobao_agent.metric_search import MetricDefinition
from taobao_agent.metric_sql import build_metric_query, validate_metric_query


UCI_CATALOG_PATH = (
    Path(__file__).resolve().parents[2] / "examples" / "uci_online_retail_metrics.json"
)
SOURCE_ID = "uci-online-retail-352"
_SOURCE_COLUMNS = (
    "InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
    "UnitPrice", "CustomerID", "Country",
)
_TABLE = "retail_lines"
_APPROVED_COLUMNS = frozenset({"invoice_date", "record_type", "line_amount_pence"})


@dataclass(frozen=True)
class ImportSummary:
    source_sha256: str
    positive_rows: int
    cancelled_rows: int
    other_rows: int
    first_date: str
    last_date: str


class UCIQueryError(Exception):
    """The local public-data query could not be safely completed."""


class UCIDateOutOfRange(UCIQueryError):
    def __init__(self, first_date: str, last_date: str) -> None:
        self.first_date = first_date
        self.last_date = last_date
        super().__init__(f"数据日期仅覆盖 {first_date} 至 {last_date}")


def _source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _line_record(row: tuple[object, ...], row_number: int) -> tuple[str, str, int]:
    invoice, _, _, quantity, invoice_date, unit_price, _, _ = row
    if (
        not isinstance(invoice, (str, int))
        or isinstance(invoice, bool)
        or not isinstance(quantity, (int, float))
        or isinstance(quantity, bool)
        or not isinstance(unit_price, (int, float))
        or isinstance(unit_price, bool)
        or not isinstance(invoice_date, (date, datetime))
    ):
        raise ValueError(f"UCI 第 {row_number} 行的查询字段不符合预期")
    if str(invoice).upper().startswith("C"):
        record_type = "cancel"
    elif quantity <= 0 or unit_price <= 0:
        record_type = "other"
    else:
        record_type = "positive"
    pence = int(
        (Decimal(str(quantity)) * Decimal(str(unit_price)) * 100).quantize(
            Decimal("1"), rounding=ROUND_HALF_UP
        )
    )
    business_date = invoice_date.date() if isinstance(invoice_date, datetime) else invoice_date
    return business_date.isoformat(), record_type, pence


def import_uci_online_retail(source_path: Path, database_path: Path) -> ImportSummary:
    """Import the official workbook without retaining customer or product details."""

    if database_path.exists():
        raise FileExistsError(f"目标数据库已存在，不会覆盖：{database_path}")
    source_sha256 = _source_hash(source_path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        if tuple(next(rows, ())) != _SOURCE_COLUMNS:
            raise ValueError("UCI 文件表头与 Online Retail 官方字段不一致")
        database_path.parent.mkdir(parents=True, exist_ok=True)
        with closing(sqlite3.connect(database_path)) as connection:
            try:
                connection.executescript(
                    "CREATE TABLE retail_lines (invoice_date TEXT NOT NULL, "
                    "record_type TEXT NOT NULL "
                    "CHECK (record_type IN ('positive', 'cancel', 'other')), "
                    "line_amount_pence INTEGER NOT NULL);"
                    "CREATE TABLE import_metadata (source_id TEXT NOT NULL, "
                    "source_sha256 TEXT NOT NULL, "
                    "first_date TEXT NOT NULL, last_date TEXT NOT NULL);"
                )
                counts = {"positive": 0, "cancel": 0, "other": 0}
                first_date: str | None = None
                last_date: str | None = None
                batch: list[tuple[str, str, int]] = []
                for row_number, row in enumerate(rows, start=2):
                    record = _line_record(row, row_number)
                    counts[record[1]] += 1
                    first_date = min(first_date, record[0]) if first_date else record[0]
                    last_date = max(last_date, record[0]) if last_date else record[0]
                    batch.append(record)
                    if len(batch) == 5000:
                        connection.executemany("INSERT INTO retail_lines VALUES (?, ?, ?)", batch)
                        batch.clear()
                if batch:
                    connection.executemany("INSERT INTO retail_lines VALUES (?, ?, ?)", batch)
                connection.execute(
                    "CREATE INDEX retail_lines_date_type "
                    "ON retail_lines(invoice_date, record_type)"
                )
                if first_date is None or last_date is None:
                    raise ValueError("UCI 文件没有可导入的交易明细")
                connection.execute(
                    "INSERT INTO import_metadata VALUES (?, ?, ?, ?)",
                    (SOURCE_ID, source_sha256, first_date, last_date),
                )
                connection.commit()
            except BaseException:
                connection.rollback()
                raise
    except BaseException:
        database_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return ImportSummary(
        source_sha256=source_sha256,
        positive_rows=counts["positive"],
        cancelled_rows=counts["cancel"],
        other_rows=counts["other"],
        first_date=first_date,
        last_date=last_date,
    )


def query_uci_metric(metric: MetricDefinition, business_date: date, database_path: Path) -> int:
    """Run one validated sum against a read-only imported UCI database."""

    calculation = metric.calculation
    if (
        metric.name != "交易明细金额"
        or metric.table != _TABLE
        or metric.field != "line_amount_pence"
        or calculation is None
        or calculation.aggregation != "sum"
        or calculation.date_field != "invoice_date"
        or calculation.condition.field != "record_type"
        or calculation.condition.value != "positive"
        or calculation.unit != "pence"
    ):
        raise UCIQueryError("UCI 指标规则不在批准范围内")
    try:
        database_uri = f"{database_path.resolve().as_uri()}?mode=ro"
        with closing(sqlite3.connect(database_uri, uri=True)) as connection:
            connection.execute("PRAGMA query_only = ON")
            metadata = connection.execute(
                "SELECT source_id, first_date, last_date FROM import_metadata"
            ).fetchone()
            if metadata is None or metadata[0] != SOURCE_ID:
                raise UCIQueryError("数据库不是本项目的 UCI 导入结果")
            if not metadata[1] <= business_date.isoformat() <= metadata[2]:
                raise UCIDateOutOfRange(metadata[1], metadata[2])
            actual_columns = frozenset(
                row[1] for row in connection.execute('PRAGMA table_info("retail_lines")')
            )
            candidate = build_metric_query(metric, business_date)
            approved = validate_metric_query(
                candidate, metric, business_date,
                {_TABLE: actual_columns & _APPROVED_COLUMNS},
            )
            row = connection.execute(approved.statement, approved.parameters).fetchone()
            if row is None or not isinstance(row[0], int):
                raise UCIQueryError("公开数据查询没有返回整数金额")
            return row[0]
    except (OSError, sqlite3.Error, ValueError) as exc:
        raise UCIQueryError("公开数据查询未通过校验或执行失败") from exc


def main() -> None:
    parser = argparse.ArgumentParser(description="导入 UCI Online Retail Excel 到本地 SQLite")
    parser.add_argument("source", type=Path, help="官方 Online Retail.xlsx 文件")
    parser.add_argument("database", type=Path, help="新 SQLite 文件，不覆盖已有文件")
    args = parser.parse_args()
    try:
        summary = import_uci_online_retail(args.source, args.database)
    except (OSError, ValueError, sqlite3.Error) as exc:
        parser.exit(2, f"导入失败：{exc}\n")
    print(json.dumps(asdict(summary), ensure_ascii=False))
